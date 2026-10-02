"""Build the Part B interim report and supporting assets from measured results."""
import csv
import json
import subprocess
import sys
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'submission/phase2'
RESULTS, IMAGES = OUT/'results', OUT/'images'
TITLE = 'Deep Learning for Energy Consumption Forecasting: A Comparative Study of Sequential Architectures'

def read_json(path): return json.loads(path.read_text(encoding='utf-8'))

def assets(df, split, metrics):
    plt.rcParams.update({'font.size':10, 'axes.spines.top':False, 'axes.spines.right':False})
    fig, ax = plt.subplots(figsize=(10,2.7))
    train_end, val_end = split['boundary_indices'][1:3]
    daily = df.Global_active_power.resample('D').mean()
    ax.plot(daily.index,daily.values,color='#263b58',lw=.65)
    for start,end,color,label in [(0,train_end,'#c2dcf1','Train 70%'),(train_end,val_end,'#f8dfb2','Validation 15%'),(val_end,len(df),'#c9e8da','Test 15%')]:
        ax.axvspan(df.index[start],df.index[end-1],color=color,alpha=.6,label=label)
    ax.set_ylabel('Daily mean power (kW)'); ax.legend(ncol=3,fontsize=9); fig.tight_layout()
    fig.savefig(IMAGES/'dataset_split.png',dpi=220); plt.close(fig)
    fig, ax = plt.subplots(figsize=(10,2.4)); ax.set_xlim(0,10); ax.set_ylim(0,2.4); ax.axis('off')
    blocks = [(0.1,.8,1.8,'UCI minute readings\nHourly means\nPast-only gap filling'),(2.35,.8,1.6,'70 / 15 / 15 split\nTrain-only scaling\n168 h x 14 inputs'),(4.5,1.65,1.6,'MLP\n256 / 128 / 64'),(4.5,.85,1.6,'1D-CNN\n32 / 64 / 128'),(4.5,.05,1.6,'LSTM\n2 layers x 64'),(7,.8,2.65,'24-hour direct forecast\nCommon observed targets\nMAE / RMSE / MAPE / R2')]
    for x,y,w,label in blocks:
        ax.add_patch(FancyBboxPatch((x,y),w,.63,boxstyle='round,pad=.07',facecolor='#eef3f8',edgecolor='#466681'))
        ax.text(x+w/2,y+.315,label,ha='center',va='center',fontsize=9)
    for a,b in [((1.95,1.1),(2.25,1.1)),((4,1.1),(4.4,1.95)),((4,1.1),(4.4,1.15)),((4,1.1),(4.4,.35)),((6.2,1.95),(6.9,1.1)),((6.2,1.15),(6.9,1.1)),((6.2,.35),(6.9,1.1))]:
        ax.annotate('',xy=b,xytext=a,arrowprops={'arrowstyle':'->','color':'#466681'})
    fig.tight_layout(); fig.savefig(IMAGES/'system_architecture.png',dpi=220); plt.close(fig)
    fig, axes = plt.subplots(1,2,figsize=(10,3.1))
    horizons = pd.read_csv(RESULTS/'horizon_mae.csv')
    for name in ['MLP','CNN1D','LSTM','Seasonal 24h']:
        axes[0].plot(horizons.hour_ahead,horizons[name],label=name,lw=1.5)
    axes[0].set(xlabel='Hours ahead',ylabel='MAE (kW)',title='Error by forecast horizon'); axes[0].legend(fontsize=8)
    y = np.load(OUT/'models/trues_test.npy')
    axes[1].plot(range(1,25),y[0],color='black',label='Observed',lw=1.8)
    for name in ['mlp','cnn1d','lstm']:
        axes[1].plot(range(1,25),np.load(OUT/f'models/preds_{name}.npy')[0],label=name.upper(),lw=1.3)
    axes[1].set(xlabel='Hours ahead',ylabel='Active power (kW)',title='First eligible test forecast'); axes[1].legend(fontsize=8)
    fig.tight_layout(); fig.savefig(IMAGES/'report_results.png',dpi=240); plt.close(fig)
    fig, axes = plt.subplots(1,3,figsize=(11,2.8))
    for ax,name in zip(axes,['mlp','cnn1d','lstm']):
        history = pd.read_csv(OUT/f'models/history_{name}.csv')
        ax.plot(history.epoch,history.train_mse_scaled,label='Train')
        ax.plot(history.epoch,history.validation_mse_scaled,label='Validation')
        ax.set(title=name.upper(),xlabel='Epoch',ylabel='Scaled MSE'); ax.legend(fontsize=8)
    fig.tight_layout(); fig.savefig(IMAGES/'learning_curves.png',dpi=220); plt.close(fig)
    metrics.to_csv(RESULTS/'report_table_metrics.csv',index=False)

def paragraph(doc,text='',size=10,bold=False,align=None):
    p = doc.add_paragraph()
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1
    r = p.add_run(text); r.font.name='Times New Roman'; r.font.size=Pt(size); r.bold=bold
    return p

def heading(doc,text):
    p=paragraph(doc,text,11,True,WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_before=Pt(6); p.paragraph_format.keep_with_next=True

def table(doc,headers,rows,widths=None,size=8.5):
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
    try: t.style='Table Grid'
    except KeyError: pass
    for i,h in enumerate(headers): t.rows[0].cells[i].text=str(h)
    for row in rows:
        for c,text in zip(t.add_row().cells,row): c.text=str(text)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        for i,c in enumerate(row.cells):
            if widths: c.width=Inches(widths[i])
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.line_spacing=1
                for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(size); r.bold=ri==0
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    return t

def figure(doc,filename,caption,width=6.9):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(1)
    p.add_run().add_picture(str(IMAGES/filename),width=Inches(width))
    paragraph(doc,caption,8,align=WD_ALIGN_PARAGRAPH.CENTER)

def report():
    if not (RESULTS/'completion.json').exists():
        raise RuntimeError('Complete the experiment before building the report')
    team=read_json(OUT/'team_details.json'); papers=read_json(OUT/'literature.json')
    audit=read_json(RESULTS/'data_audit.json'); split=read_json(RESULTS/'split_manifest.json')
    config=read_json(RESULTS/'run_config.json'); errors=read_json(RESULTS/'error_analysis.json')
    metrics=pd.read_csv(RESULTS/'model_comparison.csv')
    df=pd.read_csv(RESULTS/'hourly_causal.csv',index_col='datetime',parse_dates=True)
    assets(df,split,metrics)
    template=next((ROOT/'Context').glob('*.docx'))
    doc=Document(template)
    for child in list(doc.element.body):
        if child.tag != qn('w:sectPr'): doc.element.body.remove(child)
    section=doc.sections[0]
    section.page_width=Mm(210); section.page_height=Mm(297)
    section.top_margin=Mm(14); section.bottom_margin=Mm(14)
    section.left_margin=Mm(16); section.right_margin=Mm(16)
    for part in [section.header,section.footer]:
        for p in part.paragraphs: p.clear()
    normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(4); normal.paragraph_format.line_spacing=1
    # Cover follows the supplied title/author/affiliation structure.
    paragraph(doc,'',24)
    paragraph(doc,TITLE,24,True,WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc,'',18)
    paragraph(doc,'Interim report on\nDeep Learning Project\n[ICT-4442]',16,align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc,'Phase 2 - Part B',12,align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc,'',18)
    paragraph(doc,'Submitted By',12,True,WD_ALIGN_PARAGRAPH.CENTER)
    for member in team['members']:
        paragraph(doc,member['name']+'  |  '+member['registration'],11,align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc,'Team No.: '+team['team_number'],11,align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc,'',24)
    p=paragraph(doc,'SCHOOL OF COMPUTER ENGINEERING\nMANIPAL INSTITUTE OF TECHNOLOGY\nMANIPAL ACADEMY OF HIGHER EDUCATION',10,align=WD_ALIGN_PARAGRAPH.CENTER)
    p.runs[0].italic=True
    paragraph(doc,'OCTOBER 2026',11,align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc,'Prepared: '+team['submission_date'],10,align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc,'GitHub Repository Link: '+team['repository_url'],9,align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc,'Confirmed title: unchanged from the approved synopsis.',9,align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc,'Team verification and signatures are required before submission.',9,align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()
    heading(doc,'1. Literature Review')
    paragraph(doc,'The study predicts the next 24 hourly active-power values using the preceding seven days. Table I compares ten papers across architecture families. Published findings describe each paper\'s own protocol; their numerical scores are not directly comparable with this experiment.')
    paragraph(doc,'TABLE I. LITERATURE COMPARISON',8,True,WD_ALIGN_PARAGRAPH.CENTER)
    table(doc,['Paper (Author, Year)','Method','Dataset','Key Result','Relevance to Project'],
          [[p[k] for k in ['paper','method','dataset','result','relevance']] for p in papers],
          [1.0,1.12,1.0,1.75,2.1],8.5)
    heading(doc,'Synthesis and research gap')
    paragraph(doc,'The load-specific comparison [1] cautions that household demand is less regular than aggregated demand. Convolutional evidence [2], [5] challenges an automatic preference for recurrence, while N-BEATS and linear baselines [6], [9] show that model complexity alone does not establish forecasting quality. Attention methods [7], [8], [10] and the solar hybrid [4] offer extensions, but use different tasks and training budgets. Consequently, this project compares a plain MLP, pooled 1D-CNN and stacked LSTM with fixed inputs, splits and metrics, and includes persistence controls. Uncertainty and distribution changes identified in [3] remain limitations.')
    doc.add_page_break()
    heading(doc,'2. Dataset Acquisition & Preprocessing')
    paragraph(doc,f'The UCI Individual Household Electric Power Consumption dataset [11] contains {audit["raw_rows"]:,} minute readings from one household in Sceaux, France. The local file spans {audit["start"][:10]} to {audit["end"][:10]}; its SHA-256 checksum is recorded in results/data_audit.json. Each of the seven measurement columns has {audit["missing_raw_by_column"]["Global_active_power"]:,} missing entries. Acquisition code is supplied and the complete raw dataset was used.')
    paragraph(doc,f'Cleaning parses day/month/year timestamps, converts missing markers to null values, checks duplicates, and aggregates available minute readings to hourly means ({audit["hourly_rows"]:,} hours). Entirely missing hourly features are filled from preceding hours only. Forecast labels must have all 60 minute readings: {audit["excluded_incomplete_target_hours"]} incomplete hours are ineligible as labels. This avoids evaluating on manufactured targets; historical inputs can still contain imputed values. The first and last partial hours are also excluded as labels.')
    paragraph(doc,'Encoding uses seven electrical measurements, sine/cosine pairs for hour, weekday and month, plus a weekend flag (14 inputs). Sub-metering features retain the mean of their original minute-energy values; the target is hourly mean active power in kW, not energy in kWh. No augmentation is used. Min-max scalers are fitted on training rows only; validation/test values are not clipped. Each window stays wholly inside its split: 168 past hours predict 24 future hours, with stride one.')
    paragraph(doc,'TABLE II. SHARED CHRONOLOGICAL SPLIT',8,True,WD_ALIGN_PARAGRAPH.CENTER)
    table(doc,['Split','Hourly rows','Valid windows','Period (inclusive)'],
          [[name.title(),f'{split[name]["rows"]:,}',f'{split[name]["windows"]:,}',split[name]['start']+' to '+split[name]['end']] for name in ['train','validation','test']],
          [.85,.85,1.0,4.27],8.5)
    figure(doc,'dataset_split.png','Figure 1. Daily mean consumption and fixed chronological split.',6.8)
    heading(doc,'3. Models Implemented So Far')
    paragraph(doc,'All three proposed families are implemented and trained. The owner column follows the synopsis assignment; individual authorship must be confirmed in Section 4. Dropout is 0.2 and each network directly outputs 24 values.')
    paragraph(doc,'TABLE III. IMPLEMENTATION STATUS AND PRELIMINARY METRIC',8,True,WD_ALIGN_PARAGRAPH.CENTER)
    owners={'mlp':'Vignesh Rao','cnn1d':'Vignesh Rao','lstm':'Elton Lewis'}
    architectures={'mlp':'Flatten 168 x 14; dense 256/128/64; ReLU; dense 24.',
                   'cnn1d':'Conv filters 32/64/128, kernel 3; batch norm, ReLU, pool 2; global average; dense 64/24.',
                   'lstm':'Two recurrent layers, 64 hidden units; final state; dense 64/24.'}
    table(doc,['Model','Owner','Status / test MAE','Architecture / parameters'],
          [[r.model.upper(),owners[r.model],f'Trained; {r["MAE (kW)"]:.4f} kW',architectures[r.model]+f' {int(r.parameters):,} parameters.'] for _,r in metrics.iloc[:3].iterrows()],
          [.7,1,1.4,3.87],8.5)
    paragraph(doc,f'Common training: Adam, learning rate {config["lr"]}, scaled MSE loss, batch size {config["batch_size"]}, at most {config["epochs"]} epochs, early-stopping patience {config["patience"]}, seed {config["seed"]}; deterministic CPU execution with {config["threads"]} threads. Minimum validation loss selects the checkpoint; test scores do not guide training. No hyperparameter search was performed.',9)
    doc.add_page_break()
    heading(doc,'3. Models Implemented So Far (continued): Preliminary Results')
    paragraph(doc,f'All models use the same {split["test"]["windows"]:,} valid test windows and {split["test"]["windows"]*24:,} forecast-target pairs. Predictions are inverse-transformed to kW. MAE and RMSE measure absolute error; R2 compares squared error with the overall target mean. MAPE uses max(|actual|, 0.0001) in its denominator and can be large at low loads. Scores pool all horizons and origins; overlapping windows are correlated, so this count is not an independent sample size.')
    paragraph(doc,'TABLE IV. COMMON TEST-SET COMPARISON (SINGLE SEED)',8,True,WD_ALIGN_PARAGRAPH.CENTER)
    table(doc,['Model','MAE (kW)','RMSE (kW)','MAPE (%)','R2','Best epoch'],
          [[r.model.upper() if i<3 else r.model,f'{r["MAE (kW)"]:.4f}',f'{r["RMSE (kW)"]:.4f}',f'{r["MAPE (%)"]:.2f}',f'{r["R2 Score"]:.4f}',str(int(r.best_epoch)) if i<3 else 'N/A'] for i,r in metrics.iterrows()],
          [1.55,1.02,1.05,1.1,1.0,1.25],9)
    paragraph(doc,'Persistence repeats the last observed input-hour value. Seasonal controls reuse values from 24 or 168 hours earlier. Controls share the same target eligibility and may use causally imputed historical inputs.',9)
    figure(doc,'report_results.png','Figure 2. Horizon error and the first eligible test forecast; selection fixed independently of model performance.',6.9)
    neural=metrics.iloc[:3].sort_values('MAE (kW)'); winner=neural.iloc[0]
    seasonal=metrics.loc[metrics.model=='Seasonal 24h'].iloc[0]
    improvement=(seasonal['MAE (kW)']-winner['MAE (kW)'])/seasonal['MAE (kW)']*100
    paragraph(doc,f'{winner.model.upper()} has the lowest neural-network test MAE ({winner["MAE (kW)"]:.4f} kW) in this run, a {improvement:.1f}% reduction relative to daily seasonal persistence. This is a descriptive result from one configuration and seed, not evidence that the family is universally superior. The architectures use different parameter counts and compute; identical epochs do not imply identical optimization effort.')
    paragraph(doc,'The MLP can exploit fixed positions and calendar covariates directly. Convolution shares local filters efficiently, but global average pooling discards explicit temporal position. The LSTM carries recurrent state across 168 steps, which may help dependence modeling but is harder to optimize. These mechanisms provide plausible explanations for differences; feature and architecture ablations are needed to establish causes.')
    paragraph(doc,f'For {errors["best_neural_by_test_mae"]}, peak-hour MAE is {errors["peak_mae_kw"]:.4f} kW versus {errors["nonpeak_mae_kw"]:.4f} kW outside peaks. Peaks use the training 90th-percentile threshold ({errors["train_90th_percentile_kw"]:.4f} kW). The worst 24-hour window starts {errors["worst_sample_start"]} and has MAE {errors["worst_sample_mae_kw"]:.4f} kW. Saved worst_forecast.png supports inspection of this failure case. Household activity spikes and changes between seasons remain difficult to anticipate from historical inputs alone.')
    figure(doc,'learning_curves.png','Figure 3. Training and validation curves. Best validation checkpoints are restored for test evaluation.',6.9)
    doc.add_page_break()
    heading(doc,'4. Individual Contribution Log (to date)')
    paragraph(doc,'The matrix below records synopsis ownership and current repository evidence.',9)
    table(doc,['Member Name / Reg. No.','Task(s) Completed / Evidence to Verify','Signature'],[
        ['Vignesh Rao\n220953668','Assigned: preprocessing, MLP and CNN. Corresponding code and trained artifacts are present.','________________'],
        ['Elton Lewis\n230953572','Assigned: LSTM and evaluation. Corresponding code and trained artifacts are present.','________________']],
        [1.55,4.12,1.3],8.5)
    heading(doc,'5. Risk / Plan for Remaining Work')
    table(doc,['Target date','Remaining work / risk response'],[
        ['2-4 Oct','Verify ownership, fill team number, obtain signatures and institutional checks; confirm acceptance after the printed 25 Sep deadline.'],
        ['5-12 Oct','Run multiple seeds and validation-only tuning; compare calendar-only and reduced-history ablations. Retain this run as the interim snapshot.'],
        ['13-20 Oct','Analyze seasonal shifts, peaks and long imputation gaps. For final tuning, reserve a new holdout or use nested rolling validation; this test set is already inspected.'],
        ['21-25 Oct','Each owner writes methodology and verifies references; integrate 12+ references, final contribution statement, and 10-12 minute slides.'],
        ['26-31 Oct','Present 26-29 Oct; submit Part C report and repository by 31 Oct 2026. No proposed model remains unimplemented.']], [.8,6.17],8.5)
    heading(doc,'References')
    for i,p in enumerate(papers,1):
        paragraph(doc,f'[{i}] {p["reference"]} {p["url"]}',7.5)
    paragraph(doc,'[11] G. Hebrail and A. Berard, "Individual Household Electric Power Consumption," UCI Machine Learning Repository, 2006, doi: 10.24432/C58K54. https://doi.org/10.24432/C58K54',7.5)
    doc.core_properties.title=TITLE
    doc.core_properties.subject='ICT-4442 Phase 2 interim report, Part B'
    doc.core_properties.author='Vignesh Rao; Elton Lewis'
    doc.save(OUT/'Phase_2_Interim_Report.docx')
    with (RESULTS/'literature_review.csv').open('w',newline='',encoding='utf-8-sig') as handle:
        writer=csv.DictWriter(handle,fieldnames=list(papers[0])); writer.writeheader(); writer.writerows(papers)
    print('Created',OUT/'Phase_2_Interim_Report.docx')

if __name__ == '__main__': report()
